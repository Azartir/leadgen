"""Regression tests for CSV and XLSX export values."""
import csv

import openpyxl

from leadgen.export import write_csv, write_xlsx, write_outputs


COLUMNS = [
    ("Компания", "name"),
    ("Телефон", "phone"),
    ("Все источники", "sources"),
    ("Расхождения", "source_conflicts"),
    ("Метаданные", "metadata"),
    ("Балл", "score"),
]


def test_xlsx_handles_structured_values_and_round_trips(tmp_path):
    path = tmp_path / "leads.xlsx"
    leads = [{
        "name": "ООО «Тест»",
        "phone": "+7 900 123-45-67",
        "sources": ["osm", "2gis"],
        "source_conflicts": ["телефон: +7 900 000 <> +7 900 111"],
        "metadata": {"verified": True, "category": "Услуги"},
        "score": 52,
    }]
    write_xlsx(leads, COLUMNS, str(path), title="Российские лиды")

    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook.active
    assert sheet["A2"].value == "ООО «Тест»"
    assert sheet["B2"].value == "+7 900 123-45-67"
    assert sheet["C2"].value == "osm; 2gis"
    assert "телефон" in sheet["D2"].value
    assert '"verified": true' in sheet["E2"].value
    assert sheet["F2"].value == 52
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F2"


def test_csv_is_excel_friendly_and_preserves_structured_values(tmp_path):
    path = tmp_path / "leads.csv"
    write_csv([{
        "name": "Компания",
        "phone": "+7 900 123-45-67",
        "sources": ["osm", "2gis"],
        "source_conflicts": [],
        "metadata": {"active": True},
        "score": 9,
    }], COLUMNS, str(path))
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[0][0] == "Компания"
    assert rows[1][1] == "+7 900 123-45-67"
    assert rows[1][2] == "osm; 2gis"
    assert rows[1][3] == ""


def test_write_outputs_creates_both_files(tmp_path):
    stem = str(tmp_path / "run")
    csv_path, xlsx_path = write_outputs(
        [{"name": "Компания", "phone": "+7 900 123-45-67"}],
        COLUMNS,
        stem,
        log=lambda _: None,
    )
    assert (tmp_path / "run_crm.csv").is_file()
    assert (tmp_path / "run.xlsx").is_file()
    assert csv_path.endswith("_crm.csv")
    assert xlsx_path.endswith(".xlsx")


def test_empty_export_still_produces_valid_workbook(tmp_path):
    path = tmp_path / "empty.xlsx"
    write_xlsx([], COLUMNS, str(path))
    workbook = openpyxl.load_workbook(path)
    assert workbook.active.max_row == 1
    assert workbook.active.max_column == len(COLUMNS)
